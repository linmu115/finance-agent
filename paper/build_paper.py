"""Build editable DOCX + review HTML from manuscript.json using bundled Python."""
from __future__ import annotations
import html
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
data=json.loads((HERE/"manuscript.json").read_text(encoding="utf-8"))
updates=json.loads((HERE/"experiment-update.json").read_text(encoding="utf-8")) if (HERE/"experiment-update.json").exists() else {}
for b in data["blocks"]:
    if b.get("key") in updates:
        if isinstance(updates[b["key"]],dict):b.update(updates[b["key"]])
        else:b["text"]=updates[b["key"]]
for anchor,key in [('scoring_protocol','scorer_validation'),('model_result','model_runs'),('model_runs','model_tracks'),('model_tracks','model_costs'),('model_costs','repeatability_result'),('repeatability_result','failure_analysis'),('failure_analysis','model_score_limits'),('model_score_limits','model_report_identity'),('model_report_identity','system_acceptance_result'),('system_acceptance_result','aborted_runs')]:
    if key in updates and not any(b.get('key')==key for b in data['blocks']):
        index=next((i for i,b in enumerate(data['blocks']) if b.get('key')==anchor),None)
        if index is None:index=next(i for i,b in enumerate(data['blocks']) if b.get('key')=='model_result')
        value=updates[key]
        block={'key':key,**value} if isinstance(value,dict) else {'type':'paragraph','key':key,'text':value}
        data['blocks'].insert(index+1,block)
result=json.loads((ROOT/"benchmark/results/financial_pipeline.json").read_text(encoding="utf-8"))
for b in data["blocks"]:
    if b["type"]=="dynamic_table" and b["key"]=="pipeline_runs":
        b.update(type="table",headers=["轮次","物理页","目标事实","数值及引用通过","耗时 秒"],
                 rows=[[str(i+1),str(r["physical_pages_processed"]),str(r["total"]),f'{r["passed"]}/{r["total"]}',f'{r["elapsed_seconds"]:.3f}'] for i,r in enumerate(result["runs"])])

def create_diagram():
    im=Image.new("RGB",(1600,690),"white")
    draw=ImageDraw.Draw(im)
    fontpath="C:/Windows/Fonts/msyh.ttc"
    normal=ImageFont.truetype(fontpath,29)
    small=ImageFont.truetype(fontpath,23)
    def centered(text,box,font=normal):
        x1,y1,x2,y2=box
        b=draw.textbbox((0,0),text,font=font)
        draw.text(((x1+x2-b[2])/2,(y1+y2-b[3])/2),text,font=font,fill="black")
    def box(x,y,w,h,title,subtitle):
        draw.rounded_rectangle((x,y,x+w,y+h),radius=10,outline="#444444",width=2,fill="#fafafa")
        centered(title,(x,y+10,x+w,y+62))
        centered(subtitle,(x,y+60,x+w,y+h-8),small)
    def arrow(x1,y1,x2,y2):
        draw.line((x1,y1,x2,y2),fill="#555555",width=3)
        draw.polygon([(x2,y2),(x2-9,y2-13),(x2+9,y2-13)],fill="#555555")
    box(530,15,540,108,"对话前端","消息  附件  配置  来源核查")
    arrow(800,123,800,165)
    box(530,165,540,108,"Node 宿主与 Pi 会话","项目绑定  工具白名单  运行事件")
    arrow(800,273,800,315)
    box(530,315,540,108,"Python 金融 MCP 服务","解析  检索  精确查询  固定计算")
    draw.line((265,457,1335,457),fill="#555555",width=3)
    draw.line((800,423,800,457),fill="#555555",width=3)
    for x in (265,800,1335):arrow(x,457,x,495)
    box(45,495,440,120,"文档与 OCR","原文  表格  单元格位置")
    box(580,495,440,120,"SQLite 与文件","事实  FTS5  向量缓存  状态")
    box(1115,495,440,120,"隔离分析容器","只读输入  断网  受控产物")
    centered("证据和计算结果回传对话  宿主保管凭据  自由代码不在宿主执行",(0,635,1600,685),small)
    im.save(HERE/"architecture.png")

create_diagram()
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(2.2);section.bottom_margin=Cm(2.1)
section.left_margin=Cm(2.4);section.right_margin=Cm(2.4)
section.header_distance=Cm(.8);section.footer_distance=Cm(.9)

def style_font(style,latin,cjk,size,bold=False):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    fonts=rpr.find(qn("w:rFonts"))
    if fonts is None:fonts=OxmlElement("w:rFonts");rpr.insert(0,fonts)
    for k,v in [("ascii",latin),("hAnsi",latin),("eastAsia",cjk)]:fonts.set(qn("w:"+k),v)
    for color in rpr.findall(qn("w:color")):
        color.attrib.pop(qn("w:themeColor"),None)

style_font(doc.styles["Normal"],"Times New Roman","宋体",10.5)
base=doc.styles["Normal"].paragraph_format
base.line_spacing=1.27;base.space_after=Pt(5);base.first_line_indent=Pt(21)
style_font(doc.styles["Title"],"Times New Roman","黑体",18,True)
title_props=doc.styles["Title"].element.get_or_add_pPr()
for border in list(title_props.findall(qn("w:pBdr"))):title_props.remove(border)
doc.styles["Title"].paragraph_format.space_after=Pt(14)
doc.styles["Title"].paragraph_format.line_spacing=1.2
for name,size in [("Heading 1",13),("Heading 2",11.5)]:
    style_font(doc.styles[name],"Times New Roman","黑体",size,True)
    fmt=doc.styles[name].paragraph_format;fmt.space_before=Pt(12);fmt.space_after=Pt(6);fmt.first_line_indent=Pt(0);fmt.keep_with_next=True
style_font(doc.styles["Caption"],"Times New Roman","宋体",9)
doc.styles["Caption"].paragraph_format.first_line_indent=Pt(0)
doc.styles["Caption"].paragraph_format.space_before=Pt(6)
doc.styles["Caption"].paragraph_format.space_after=Pt(6)
doc.styles["Caption"].paragraph_format.keep_with_next=True
title=doc.add_paragraph(data["title"].replace("的财报","的\n财报"),"Title")
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.first_line_indent=Pt(0)
for text in (data["authors"],data["version"]):
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(8)
    for run in p.runs:run.font.size=Pt(10)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
footer._p.append(field)

def cell_borders(cell):
    pr=cell._tc.get_or_add_tcPr()
    borders=OxmlElement("w:tcBorders")
    for name in ["top","left","bottom","right","insideH","insideV"]:
        tag=OxmlElement("w:"+name);tag.set(qn("w:val"),"single");tag.set(qn("w:sz"),"4");tag.set(qn("w:color"),"D9D9D9");borders.append(tag)
    pr.append(borders)
    margins=OxmlElement("w:tcMar")
    for name,value in [("top","90"),("bottom","90"),("left","100"),("right","100")]:
        tag=OxmlElement("w:"+name);tag.set(qn("w:w"),value);tag.set(qn("w:type"),"dxa");margins.append(tag)
    pr.append(margins)

htmlparts=[]
for block in data["blocks"]:
    kind=block["type"]
    if kind=="heading":
        doc.add_paragraph(block["text"],"Heading "+str(block["level"]))
        htmlparts.append(f'<h{block["level"]+1}>{html.escape(block["text"])}</h{block["level"]+1}>')
    elif kind in {"paragraph","reference"}:
        p=doc.add_paragraph(block["text"])
        if kind=="reference":
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(5)
            for run in p.runs:run.font.size=Pt(8.5)
        p.paragraph_format.widow_control=True
        htmlparts.append('<p'+(' class="reference"' if kind=="reference" else '')+'>'+html.escape(block["text"])+'</p>')
    elif kind=="figure":
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(HERE/"architecture.png"),width=Cm(16.2))
        cap=doc.add_paragraph(block["caption"],"Caption");cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next=False
        htmlparts.append('<figure><img src="architecture.png" alt="前端经Node和Pi调用金融MCP，连接文档、SQLite与隔离分析容器"><figcaption>'+html.escape(block["caption"])+'</figcaption></figure>')
    elif kind=="table":
        doc.add_paragraph(block["caption"],"Caption")
        columns=len(block["headers"])
        table=doc.add_table(rows=1,cols=columns)
        table.autofit=False
        widths=([3.0,6.6,6.6] if columns==3 else [6.6]+[3.2]*(columns-1) if columns==4 else [2.1,2.4,2.7,5.0,4.0])
        widths=[w*16.2/sum(widths) for w in widths]
        for i,width in enumerate(widths):table.columns[i].width=Cm(width)
        for i,head in enumerate(block["headers"]):table.rows[0].cells[i].text=head
        for row in block["rows"]:
            cells=table.add_row().cells
            for i,text in enumerate(row):cells[i].text=str(text)
        repeat=OxmlElement("w:tblHeader");table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index,row in enumerate(table.rows):
            cant=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(cant)
            for col,cell in enumerate(row.cells):
                cell.width=Cm(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell_borders(cell)
                if index==0:
                    shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EEEEEE");cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
                    if len(table.rows)<=10:p.paragraph_format.keep_with_next=index<len(table.rows)-1
                    if columns>=4:p.alignment=WD_ALIGN_PARAGRAPH.CENTER if col>0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:run.font.size=Pt(9);run.font.bold=index==0
        htmlparts.append('<figure><figcaption>'+html.escape(block["caption"])+'</figcaption><table><thead><tr>'+''.join('<th>'+html.escape(v)+'</th>' for v in block["headers"])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in row)+'</tr>' for row in block["rows"])+'</tbody></table></figure>')

doc.core_properties.title=data["title"];doc.core_properties.author="";doc.core_properties.subject="金融人工智能系统工程论文初版"
out=HERE/"financial_analysis_paper_v0.1.docx";doc.save(out)
css="""@page{size:A4;margin:22mm 24mm}body{max-width:850px;margin:55px auto;padding:0 25px;color:#111;font:16px/1.9 'SimSun','Songti SC',serif}h1{text-align:center;font:600 28px/1.5 'Microsoft YaHei',sans-serif;margin:0 0 24px}header p{text-align:center;text-indent:0;color:#555;font-size:13px;margin:5px}h2{font:600 21px/1.5 'Microsoft YaHei',sans-serif;margin:30px 0 12px}h3{font:600 17px/1.5 'Microsoft YaHei',sans-serif;margin:20px 0 9px}p{text-indent:2em;text-align:justify;margin:10px 0}figure{margin:20px 0}img{width:100%;height:auto}figcaption{text-align:center;font-size:13px;margin:8px}table{border-collapse:collapse;width:100%;font-size:13px;line-height:1.65}th,td{border:1px solid #d9d9d9;padding:9px 10px;vertical-align:middle}th{background:#eee}p.reference{font-size:12px;line-height:1.6;text-indent:0;overflow-wrap:anywhere}@media print{body{padding:0;margin:0;font-size:10.5pt;line-height:1.6}h2,h3,figcaption{break-after:avoid}figure{break-inside:avoid}header p{font-size:10pt}}"""
(HERE/"financial_analysis_paper_v0.1.html").write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(data["title"])+'</title><style>'+css+'</style><body><header><h1>'+html.escape(data["title"])+'</h1><p>'+html.escape(data["authors"])+'</p><p>'+html.escape(data["version"])+'</p></header>'+''.join(htmlparts)+'</body></html>',encoding="utf-8")
(HERE/"content_snapshot.json").write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding="utf-8")
print(out)

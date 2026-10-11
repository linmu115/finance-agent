"""Replay only the authorized Yili snapshot; never open or write the live DB."""
from pathlib import Path
from types import SimpleNamespace
import hashlib
import json
import subprocess
import time
import sys
import fitz
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT))
from service.report_package import prepare,finish

OUT=Path(__file__).resolve().parent
PROJECT='374beac2-a970-4972-b210-e91018bbd6e6'
SOURCE=ROOT/'data/active/report_packages'/PROJECT/'0bc31432472e43c396dd5c0943d9044a'
raw=(SOURCE/'snapshot.json').read_bytes();snapshot=json.loads(raw)
assert snapshot['conversation_id']=='b2858f3c-98b9-4eea-9f18-1ac9a7117ac4'
def project_check(project):
    assert project==PROJECT
def artifact(project,run,name):
    project_check(project)
    assert run=='24929508fdb84826a8643cc3653e9bd5' and name=='cfo_bridge.png'
    return SOURCE/'assets/1-cfo_bridge.png'
service=SimpleNamespace(store=SimpleNamespace(root=OUT/'replay',project=project_check),artifact_path=artifact)
start=time.perf_counter();prepared=prepare(service,PROJECT,snapshot);prepared_at=time.perf_counter()
folder=service.store.root/'report_packages'/PROJECT/prepared['package_id']
subprocess.run(['node',str(OUT/'render-markdown.mjs'),str(folder/'report.md'),str(OUT/'replay.html')],check=True,timeout=20,cwd=ROOT)
html_at=time.perf_counter();result=finish(service,PROJECT,prepared['package_id'],(OUT/'replay.html').read_text(encoding='utf-8'));finished=time.perf_counter()
assert json.loads((folder/'snapshot.json').read_text(encoding='utf-8'))==snapshot
assert hashlib.sha256((SOURCE/'snapshot.json').read_bytes()).digest()==hashlib.sha256(raw).digest()
old=load_workbook(SOURCE/'workpapers.xlsx');new=load_workbook(folder/'workpapers.xlsx')
assert old.sheetnames==new.sheetnames
count=0
for a,b in zip(old,new):
    assert (a.max_row,a.max_column)==(b.max_row,b.max_column)
    for row_a,row_b in zip(a,b):
        for x,y in zip(row_a,row_b):
            assert (x.value,x.data_type,x.comment.text if x.comment else None)==(y.value,y.data_type,y.comment.text if y.comment else None),(a.title,x.coordinate)
            count+=1
old.close();new.close()
images=[];overflow=[]
with fitz.open(folder/'report.pdf') as pdf:
    anomaly=[]
    for i,page in enumerate(pdf):
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for span in line['spans']:
                    if span['text'].strip()=='异常信号' and span['size']>=16:anomaly.append(i)
                    if span['bbox'][0]<20 or span['bbox'][2]>page.rect.width-20:overflow.append({'page':i+1,'bbox':span['bbox'],'text':span['text'][:60]})
    assert anomaly,'Missing anomaly chapter'
    for i in dict.fromkeys([0,1,anomaly[0]]):
        target=OUT/(f'report-anomaly-page-{i+1:02}.png' if i==anomaly[0] else f'report-page-{i+1:02}.png')
        pdf[i].get_pixmap(matrix=fitz.Matrix(1.75,1.75),alpha=False).save(target);images.append(target.relative_to(ROOT).as_posix())
    locator_page=next((i for i in range(anomaly[0]+1,len(pdf)) if '定位' in pdf[i].get_text() and '页等' in pdf[i].get_text()),None)
    if locator_page is not None:
        target=OUT/f'report-locators-page-{locator_page+1:02}.png'
        pdf[locator_page].get_pixmap(matrix=fitz.Matrix(1.75,1.75),alpha=False).save(target);images.append(target.relative_to(ROOT).as_posix())
    pages=len(pdf)
evidence={'snapshot_source':(SOURCE/'snapshot.json').relative_to(ROOT).as_posix(),'source_sha256':hashlib.sha256(raw).hexdigest(),
 'conversation_id':snapshot['conversation_id'],'project_id':PROJECT,'package_id':prepared['package_id'],
 'seconds':{'prepare':round(prepared_at-start,3),'markdown_to_html':round(html_at-prepared_at,3),'finish_pdf_and_zip':round(finished-html_at,3),'total_export':round(finished-start,3)},
 'pdf_pages':pages,'pdf_bytes':(folder/'report.pdf').stat().st_size,'markdown_bytes':(folder/'report.md').stat().st_size,
 'snapshot_equal':True,'workbook_cells_compared':count,'workbook_values_types_comments_equal':True,
 'pdf':(folder/'report.pdf').relative_to(ROOT).as_posix(),'images':images,'anomaly_pages':[i+1 for i in anomaly],'locator_summary_page':locator_page+1 if locator_page is not None else None,'horizontal_overflow':overflow}
(OUT/'real-replay.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(evidence,ensure_ascii=False))

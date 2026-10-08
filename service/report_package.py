"""Shared delivery infrastructure: immutable conversation snapshots, no inference calls."""
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import re
import shutil
import uuid
import zipfile

from .analysis_report import render_table, column_label
from openpyxl import Workbook
import fitz

CHAPTERS = ['结论', '关键数据', '业绩变化归因', '特殊损益与口径', '异常信号', '估值', '假设、风险与局限', '来源']
MISSING = ['未提供分析结论。','未保存可复用的关键经营数据表。','未提供业绩归因分析或底稿。','未提供特殊损益与口径分析或底稿。','尚未提供异常信号分析，不能据此判断没有异常。','未提供估值分析。','未提供假设、风险与局限说明。','未提供可定位来源；已留存的输入与计算链见底稿工作簿。']
# Stable metadata extension point for N2/N4: report_section can be a chapter or key below.
SECTION_KEYS = {'conclusion':0,'key_data':1,'variance_attribution':2,'basis_changes':3,'anomaly_signals':4,'valuation':5,'limitations':6,'sources':7}
TABLE_SECTIONS = {
    'bridge':2,'indirect_bridge':2,'indirect_contributions':2,'contributions':2,'drivers_evidence':2,
    'deducted_profit':3,'nonrecurring_items':3,'basis_changes':3,'restatement':3,'recurring_specials':3,
    'signals':4,'anomaly_signals':4,'comparability':3,'report_columns':7,'missing':6,'unrecognized_rows':6,
    'valuation':5,'dcf_summary':5,'dcf_forecast':5,'dcf_sensitivity':5,'dcf_update_contributions':5,
}
SHEETS = {1:'指标',2:'利润桥',3:'特殊损益与口径',4:'信号',5:'估值',6:'缺项与局限',7:'来源'}
TABLE_TITLES = {'income':'利润表','profit_attribution':'利润归属','ratios':'利润率','growth':'同比增长','cash':'现金流','working_capital':'营运资本','cash_quality':'现金转化与周转','observations':'数值关系','roe':'净资产收益率','bridge':'现金桥','indirect_bridge':'间接法现金桥','indirect_contributions':'现金桥贡献','deducted_profit':'扣非利润','nonrecurring_items':'非经常性损益','quarterly':'季度数据','quarterly_reconciliation':'季度勾稽','missing':'缺失资料','report_columns':'比较列来源','signals':'异常信号','contributions':'变化贡献','drivers_evidence':'驱动因素原文'}
TABLE_TITLES.update(quarters='单季数据', cumulatives='累计数据', changes='期间变化', reconciliation='期间勾稽', markers='口径变化原文', adjustments='重述调整', ratio_impacts='比率影响', repeated_nonrecurring='多期非经常项目', summary='口径识别结论', unrecognized='未识别资料')
TABLE_TITLES.update(dcf_summary='DCF 估值结果',dcf_forecast='FCFF 预测与折现',dcf_sensitivity='DCF 敏感性')

def table_title(name,table):
    return table.get('title') or table.get('label') or TABLE_TITLES.get(table.get('original_table',name),name)

def section_for(name,table):
    explicit=table.get('report_section')
    if explicit in CHAPTERS:return CHAPTERS.index(explicit)
    if explicit in SECTION_KEYS:return SECTION_KEYS[explicit]
    return TABLE_SECTIONS.get(table.get('original_table',name),1)

def presentation_table(table):
    """Keep technical row IDs in the workbook; they are not financial report columns."""
    columns={k:{**c,'label':c.get('label') or k} for k,c in table['columns'].items() if c.get('label')!='行键'}
    return {**table,'columns':columns}

def prose_sections(body):
    result={i:[] for i in range(8)};section=0;fenced=False
    patterns=[(6,r'假设|风险|适用范围|局限|限制'),(3,r'特殊损益|非经常|扣非|口径|重述'),(4,r'异常|信号'),(2,r'归因|驱动|利润桥|贡献'),(5,r'估值|目标价|DCF'),(1,r'关键数据|指标|财务数据'),(7,r'来源|引用'),(0,r'结论|摘要')]
    for line in body.splitlines():
        if line.lstrip().startswith('```'):fenced=not fenced
        if not fenced and re.match(r'^#{1,6}\s',line):
            section=next((i for i,p in patterns if re.search(p,line,re.I)),section)
            line='### '+re.sub(r'^#+\s*','',line)
        result[section].append(line)
    return result

def value_for_excel(value):
    if not isinstance(value,(str,int,float,Decimal)) or isinstance(value,bool):return value
    try:d=Decimal(str(value))
    except InvalidOperation:return str(value)
    if not d.is_finite():return str(value)
    # Excel holds 15 significant digits. Longer values remain exact text, never rounded on export.
    return float(d) if len(d.normalize().as_tuple().digits)<=15 else str(value)

def workbook(snapshot,output):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    book=Workbook();book.remove(book.active)
    for name in ['指标','利润桥','贡献','特殊损益与口径','信号','估值','缺项与局限','来源']:book.create_sheet(name)
    heading_rows={s.title:{1,2} for s in book}
    for sheet in book:
        sheet.append([snapshot['title']+' · '+sheet.title]);sheet.append(['按回答版本保存；数值保持原精度，超过 Excel 15 位有效数字的值以文本保留。'])
        sheet.freeze_panes='B4';sheet.sheet_view.showGridLines=False
    sources=book['来源'];sources.append(['回答版本','表名','行键','列名','来源链 / 缺失原因'])
    heading_rows['来源'].add(3)
    for n,turn in enumerate(snapshot['turns'],1):
        for ref in turn['sources']:sources.append([n,'回答','','',json.dumps(ref,ensure_ascii=False)])
        for name,t in turn['tables'].items():
            section=section_for(name,t)
            sheet=book['贡献' if name.endswith('contributions') else SHEETS.get(section,'指标')]
            fields=[(k,unit) for k,c in t['columns'].items() for unit in ([False,True] if c.get('per_cell_unit') else [False])]
            sheet.append([]);sheet.append([f'回答 {n} · {table_title(name,t)}']);heading_rows[sheet.title].update([sheet.max_row,sheet.max_row+1])
            sheet.append([(t['columns'][k].get('label') or k)+'单位' if unit else column_label({**t['columns'][k],'label':t['columns'][k].get('label') or k}) for k,unit in fields])
            for i,row in enumerate(t['rows']):
                units=(t.get('cell_units') or [{} for _ in t['rows']])[i]
                sheet.append([' '.join(str(units.get(k,{}).get(part,'')) for part in ('currency','unit')).strip() if unit else value_for_excel(row.get(k)) for k,unit in fields])
                excel_row=sheet.max_row
                for j,(k,is_unit) in enumerate(fields,1):
                    if is_unit:continue
                    cell=sheet.cell(excel_row,j)
                    if isinstance(cell.value,str):cell.data_type='s' # Source text cannot become an Excel formula.
                    cell.number_format='General'
                    ev=(t.get('evidence') or [{} for _ in t['rows']])[i].get(k)
                    unit={**t['columns'][k],**(t.get('cell_units') or [{} for _ in t['rows']])[i].get(k,{})}
                    note={'unit':unit.get('unit'),'currency':unit.get('currency'),'source':ev}
                    cell.comment=Comment(json.dumps(note,ensure_ascii=False),'Finance Agent')
                    if ev:sources.append([n,name,(t.get('row_keys') or list(map(str,range(len(t['rows'])))))[i],k,json.dumps(ev,ensure_ascii=False)])
            for note in t.get('display_notes',[]):sheet.append([note+'（本工作簿保留精确值）'])
    for sheet in book:
        for row in sheet:
            for c in row:
                c.font=Font(name='Microsoft YaHei',size=10,color='172B4D')
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,str):c.data_type='s'
            if row[0].row in heading_rows[sheet.title]:
                for c in row:c.fill=PatternFill('solid',fgColor='E8EEF5');c.font=Font(name='Microsoft YaHei',size=11,bold=True,color='17365D')
        for j in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(j)].width=32 if j==1 else 24
        if sheet.title=='来源':sheet.column_dimensions['E'].width=100
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(5,sheet.max_column))
        sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max(5,sheet.max_column))
        sheet.row_dimensions[1].height=24;sheet.row_dimensions[2].height=32
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    book.save(output)

def prepare(service,project,snapshot):
    service.store.project(project)
    if snapshot['project_id']!=project:raise ValueError('报告包项目不匹配')
    package_id=uuid.uuid4().hex;folder=service.store.root/'report_packages'/project/package_id;folder.mkdir(parents=True)
    sections={i:[] for i in range(8)}
    for n,turn in enumerate(snapshot['turns'],1):
        label=f'回答 {n} · {turn["created_at"]}'
        body=turn['body']
        for t in turn['tables'].values():
            body=body.replace('【事实/数据】\n\n'+render_table(t),'').replace(render_table(t),'')
        # CommonMark does not close punctuation-ending bold spans before adjacent CJK text.
        body=''.join(part if i%2 else re.sub(r'\*\*([^*\n]+)\*\*(?=\S)',r'**\1** ',part)
                     for i,part in enumerate(re.split(r'(```[\s\S]*?```|`[^`\n]*`)',body)))
        for artifact in turn.get('artifacts',[]):
            match=re.fullmatch(r'/api/projects/([^/]+)/artifacts/([^/]+)/([^/?]+)',artifact.get('url',''))
            if not match or not artifact.get('mimeType','').startswith('image/'):continue
            if match[1]!=project:raise ValueError('报告图片跨项目')
            from urllib.parse import unquote
            original=service._runner().artifact_path(match[2],unquote(match[3]),project_id=project)
            asset=f'assets/{n}-{original.name}';(folder/'assets').mkdir(exist_ok=True);shutil.copyfile(original,folder/asset)
            body=body.replace(artifact['url'],asset)
        for i,lines in prose_sections(body).items():
            text='\n'.join(lines).strip()
            if text:sections[i].append('### '+label+'\n\n'+text)
        for name,t in turn['tables'].items():
            if t['rows']:sections[section_for(name,t)].append('### '+label+' · '+table_title(name,t)+'\n\n【事实/数据】\n\n'+render_table(presentation_table(t)))
        if turn['sources']:sections[7].append('### '+label+'\n\n'+'\n'.join('- '+json.dumps(ref,ensure_ascii=False) for ref in turn['sources']))
        # Full per-cell provenance is in the workbook and snapshot; retain a readable locator list here.
        refs={}
        def visit(value):
            if isinstance(value,dict):
                if value.get('document_id') or value.get('source_id'):
                    ref={k:v for k,v in value.items() if k in ('document_id','page','table_id','row','column','cell','record_id','source_id','url')}
                    refs[json.dumps(ref,ensure_ascii=False)]=ref
                for v in value.values():visit(v)
            elif isinstance(value,list):
                for v in value:visit(v)
        visit(turn['tables'])
        if refs:sections[7].append('### '+label+' · 底稿来源\n\n'+'\n'.join('- '+k for k in refs))
    markdown='# 财报分析报告\n\n'+snapshot['title']+'\n\n'+ '\n\n'.join('## '+chapter+'\n\n'+('\n\n'.join(sections[i]) if sections[i] else MISSING[i]) for i,chapter in enumerate(CHAPTERS) if i!=5 or sections[5])
    (folder/'report.md').write_text(markdown,encoding='utf8')
    (folder/'snapshot.json').write_text(json.dumps(snapshot,ensure_ascii=False,indent=2),encoding='utf8')
    workbook(snapshot,folder/'workpapers.xlsx')
    return {'package_id':package_id,'markdown':markdown}

def finish(service,project,package_id,html):
    service.store.project(project)
    if not re.fullmatch('[a-f0-9]{32}',package_id):raise ValueError('报告包标识无效')
    folder=service.store.root/'report_packages'/project/package_id
    if not (folder/'snapshot.json').is_file():raise ValueError('报告包不存在')
    css='body{font-family:sans-serif;font-size:10pt;line-height:1.5;color:#172b4d} h1{font-size:24pt} h2{font-size:17pt;color:#17365d} h3{font-size:12pt} table{border-collapse:collapse;width:100%;font-size:8pt;margin:10pt 0} th,td{border:0.5pt solid #b8c8d8;padding:4pt;overflow-wrap:anywhere} img{max-width:100%} pre{white-space:pre-wrap;font-size:8pt} a{color:#235789}'
    story=fitz.Story(html=html,user_css=css,archive=fitz.Archive(str(folder)))
    target=folder/'report.pdf';writer=fitz.DocumentWriter(str(target))
    page=fitz.paper_rect('a4');rect=page+(36,42,-36,-42)
    def rectfn(rect_num,filled):return page,rect,None
    story.write(writer,rectfn);writer.close()
    with fitz.open(target) as doc:
        for i,p in enumerate(doc):p.insert_text((page.width-75,page.height-22),f'{i+1} / {len(doc)}',fontsize=8,color=(.35,.4,.45))
        doc.saveIncr()
    archive=folder/'report_package.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(folder.rglob('*')):
            if p.is_file() and p!=archive:z.write(p,p.relative_to(folder).as_posix())
    return {'package_id':package_id,'relative_path':str(archive.relative_to(service.store.root)),'files':['report.md','report.pdf','workpapers.xlsx','snapshot.json']}

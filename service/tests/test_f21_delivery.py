import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from service.delivery_text import direction_issues, xlsx_direction_issues
from service.core import FinanceService
from service.evaluation_isolation import evaluation_profile
from service.tests.test_evaluation_isolation import scoped

FIXTURES=Path(__file__).resolve().parents[2]/'tests/fixtures/f21'


class F21DirectionTests(unittest.TestCase):
    def test_review_qualified_xlsx_false_positives_and_notices(self):
        from openpyxl import load_workbook
        sources=json.loads((FIXTURES/'xlsx-correct-sources.json').read_text('utf8'))
        cells=0
        for source in sources:
            path=FIXTURES/('correct-'+source['case']+'.xlsx')
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),source['sha256'])
            self.assertEqual(xlsx_direction_issues(path),[])
            workbook=load_workbook(path,read_only=True,data_only=False)
            try:cells+=sum(isinstance(c.value,str) and c.data_type!='f' for sheet in workbook for row in sheet for c in row)
            finally:workbook.close()
        self.assertEqual(cells,92)

    def test_historical_correct_direction_lines(self):
        rows=json.loads((FIXTURES/'history-correct.json').read_text('utf8'))
        self.assertEqual(len(rows),230)
        for row in rows:
            with self.subTest(source=row['source'],line=row['line']):
                self.assertEqual(direction_issues(row['text']),[])

    def test_real_failures_and_unchanged_workbook(self):
        g1=direction_issues((FIXTURES/'G1_X3.md').read_text('utf8'),'G1_X3.md')
        self.assertTrue(any(i['kind']=='direction_order' and '8,463.7降至10,284.3' in i['text'] for i in g1))
        g7=direction_issues((FIXTURES/'G1_X7.md').read_text('utf8'),'G1_X7.md')
        self.assertTrue(any(i['kind']=='direction_relation' and '方向相反' in i['text'] for i in g7))
        path=FIXTURES/'F1_X7.xlsx';before=hashlib.sha256(path.read_bytes()).hexdigest()
        issues=xlsx_direction_issues(path)
        found=[i for i in issues if i['sheet']=='Calculations' and i['cell']=='F10']
        self.assertEqual(len(found),1);self.assertEqual(found[0]['severity'],'notice')
        self.assertIn('需核对',found[0]['message']);self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)

    def test_order_chinese_english_signed_values_and_units(self):
        for text in ['由 8,463.7 降至 10,284.3','从1亿元降低至2亿元','由 -3 降至 -1',
                     '由 2 减至 3','由 3 增至 2','由3提高到1','由 3 升至 3',
                     'Revenue decreased from 8,463.7 to 10,284.3 million.',
                     'Revenue rose from 3 to 2','Costs from 3 increased to 2',
                     '由 **1** 上升至 **0**','从1亿元上升至200万元']:
            with self.subTest(text=text):self.assertTrue(any(i['kind']=='direction_order' for i in direction_issues(text)))
        for text in ['由8,463.7升至10,284.3','从 -5 上升至 -3','由 -1 降至 -3',
                     '由3降低至1','由3亿元降至2百万元','由2百万元升至1亿元',
                     'Profit increased from -5 to -3 million','Profit fell from 3 to 2',
                     '从3%降低至2%','从3亿元降至2%','`由3增至1`']:
            with self.subTest(text=text):self.assertFalse(direction_issues(text))

    def test_relations_sentence_and_table_row_scopes(self):
        for text in ['毛利率 -3.04pp，费用率 -3.04pp，方向相反。',
                     'A +2 B +3 反向','A -2 B +3 同向','A +2 B -3 同方向',
                     'A -3pp and B -2pp are opposite','A +3 and B -2 move in the same direction',
                     '|A|-2|B|-3|方向不同|']:
            with self.subTest(text=text):
                issues=direction_issues(text);self.assertEqual(len(issues),1);self.assertNotIn('severity',issues[0])
        for text in ['A +2，B -3，方向相反。','A -2，B -3，同向。','A +2 and B +3 same direction',
                     'A -2 and B +3 opposite','A +2 B +3 并非方向相反','A -2 B +3 not same direction']:
            with self.subTest(text=text):self.assertEqual(direction_issues(text),[])
        for text in ['毛利率和费用率 opposite','A +2 B -3。其他事项方向相反。','A +2 B -3; other movements opposite',
                     'A +2 B +3 C -4 opposite','A +0 B -3 opposite','|费用率|-3.04pp|同一分母、方向相反|']:
            with self.subTest(text=text):
                issues=direction_issues(text);self.assertEqual(len(issues),1);self.assertEqual(issues[0]['severity'],'notice')
        csv='A,B,说明\n-3,-2,"两项变化，方向相反"'
        issues=direction_issues(csv,'a.csv',True);self.assertEqual(len(issues),1)
        self.assertEqual(issues[0]['line'],2);self.assertNotIn('severity',issues[0])

    def test_xlsx_strings_relation_across_cells_and_locations(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'test.xlsx';w=Workbook();s=w.active;s.title='变动'
            s.append(['A -3pp','B -2pp','方向相反']);s.append(['由2降至3']);s.append(['方向不同'])
            s.append(['=CONCAT("由2降至3")']);w.save(path);w.close()
            issues=xlsx_direction_issues(path,'test.xlsx')
            self.assertEqual([(i['kind'],i['sheet'],i['cell']) for i in issues],
                [('direction_relation','变动','C1'),('direction_order','变动','A2'),('direction_relation','变动','A3')])
            self.assertEqual(issues[-1]['severity'],'notice')

    def test_xlsx_one_unresolved_notice_per_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'test.xlsx';w=Workbook();w.active.append(['方向相反','opposite']);w.save(path);w.close()
            issues=xlsx_direction_issues(path)
            self.assertEqual(len(issues),1);self.assertEqual(issues[0]['severity'],'notice');self.assertEqual(issues[0]['cell'],'A1')

    def test_render_scopes_xlsx_and_separates_nonblocking_notices(self):
        with tempfile.TemporaryDirectory() as tmp:
            service=FinanceService(Path(tmp));service.create_project('p','Fixture')
            class Runner:
                def artifact_path(self,rid,name,project_id):return FIXTURES/name
            service.sandbox=Runner()
            try:
                with scoped('own'):
                    service._save_analysis_state('p','r',{'status':'succeeded'})
                    result=service.render_answer('p','方向相反',{},degrade=True,files=[{'run_id':'r','name':'F1_X7.xlsx'}])['data']
                    self.assertFalse(result['issues']);self.assertEqual(len(result['notices']),2)
                    self.assertEqual(result['notices'][1]['cell'],'F10')
                with scoped('foreign'):
                    foreign=service.render_answer('p','',{},degrade=True,files=[{'run_id':'r','name':'F1_X7.xlsx'}])['data']
                    absent=service.render_answer('p','',{},degrade=True,files=[{'run_id':'absent','name':'F1_X7.xlsx'}])['data']
                    self.assertEqual(foreign,absent);self.assertNotIn('opposite',json.dumps(foreign))
                token=evaluation_profile.set('original_pi')
                try:
                    with patch('service.delivery_text.direction_issues',side_effect=AssertionError('original_pi linted')):
                        r=service.render_answer('p','由1降至2 {{bad}}',{},files=[{'run_id':'r','name':'F1_X7.xlsx'}])['data']
                    self.assertEqual(r['body'],'由1降至2 {{bad}}');self.assertFalse(r['issues'])
                finally:evaluation_profile.reset(token)
            finally:
                for name in ('pool','analysis_pool','embedding_pool','processing_pool'):getattr(service,name).shutdown(wait=True)


if __name__=='__main__':unittest.main()

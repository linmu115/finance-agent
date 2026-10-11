import json
from pathlib import Path
import tempfile
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor
from html import escape
from threading import Event
import unittest
from unittest.mock import patch
from types import SimpleNamespace
from openpyxl import load_workbook
import fitz
from service.analysis_report import compile_report
from service.report_package import prepare, finish, section_for, pdf_html, readable_inline_locators

class ReportPackageTests(unittest.TestCase):
    def test_inline_provenance_is_readable_without_changing_other_json(self):
        locator=json.dumps([{'document_id':'doc','page':61,'row':2,'text':'原文|值'},{'page':64}],ensure_ascii=False).replace('|','\\|')
        text='关系方向需核对：'+locator+'；数值 {"value": -3}；普通 [说明]。'
        shown=readable_inline_locators(text)
        self.assertIn('第61、64页',shown);self.assertNotIn(locator,shown)
        self.assertIn('数值 {"value": -3}；普通 [说明]。',shown)

    def snapshot(self, tables=None):
        return {'project_id':'p','conversation_id':'same-conversation','title':'长定位报告','turns':[
            {'id':'m','created_at':'2026-10-12','body':'现金流分析。','sources':[],'tables':tables or {}}]}

    def test_long_locator_and_tall_html_cells_finish_with_exact_workbook(self):
        locator=json.dumps([{'document_id':'doc','page':61+i%5,'row':i,'text':'原文|证据'*5} for i in range(190)],ensure_ascii=False)
        self.assertGreater(len(locator),15000)
        report=compile_report({'tables':{'signals':{'columns':{'key':{},'evidence':{'label':'文档/页/行/坐标'}},
            'rows':[{'key':'signal','evidence':locator}]}}})
        snapshot=self.snapshot(report['tables'])
        with tempfile.TemporaryDirectory() as tmp:
            svc=SimpleNamespace(store=SimpleNamespace(root=Path(tmp),project=lambda p:None))
            prepared=prepare(svc,'p',snapshot);folder=Path(tmp)/'report_packages/p'/prepared['package_id']
            self.assertNotIn(locator,prepared['markdown']);self.assertIn('第61、62、63、64、65页',prepared['markdown'])
            self.assertIn('workpapers.xlsx / snapshot.json',prepared['markdown'])
            self.assertEqual(json.loads((folder/'snapshot.json').read_text('utf8')),snapshot)
            book=load_workbook(folder/'workpapers.xlsx')
            row=next(row for row in book['信号'] if row[0].value=='signal');self.assertEqual(row[1].value,locator);book.close()
            # finish also bounds old / independently supplied HTML, including
            # nested block markup and wide rows, not only new prepare output.
            html='<h1>异常信号</h1><table><tr>'+''.join('<th>列'+str(i)+'</th>' for i in range(12))+'</tr><tr>'
            html+=''.join('<td>'+escape(locator)+'</td>' for _ in range(11))+'<td>'+'<p>长段落<br/>来源</p>'*1000+'</td></tr></table>'
            bounded=pdf_html(html);self.assertLess(len(bounded),6000);self.assertIn('列组',bounded)
            start=time.monotonic();finish(svc,'p',prepared['package_id'],html)
            self.assertLess(time.monotonic()-start,30)
            with fitz.open(folder/'report.pdf') as pdf:self.assertGreater(len(pdf),0)
            self.assertFalse((folder/'.render.html').exists());self.assertFalse((folder/'.report.pdf').exists())

    def test_same_conversation_render_is_exclusive_even_for_different_packages(self):
        with tempfile.TemporaryDirectory() as tmp:
            svc=SimpleNamespace(store=SimpleNamespace(root=Path(tmp),project=lambda p:None))
            first=prepare(svc,'p',self.snapshot());second=prepare(svc,'p',self.snapshot())
            entered=Event();release=Event();calls=[]
            def worker(args,**kwargs):
                calls.append(args);entered.set();self.assertTrue(release.wait(5))
                with fitz.open() as doc:
                    doc.new_page();doc.save(Path(args[-1])/'.report.pdf')
                return SimpleNamespace(returncode=0)
            with patch('service.report_package.subprocess.run',side_effect=worker),ThreadPoolExecutor() as pool:
                running=pool.submit(finish,svc,'p',first['package_id'],'<p>test</p>')
                try:
                    self.assertTrue(entered.wait(5))
                    with self.assertRaisesRegex(ValueError,'仍在导出'):finish(svc,'p',second['package_id'],'<p>test</p>')
                    self.assertEqual(len(calls),1)
                finally:release.set()
                running.result(timeout=5)

    def test_long_table_finishes_without_repeating_or_losing_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            svc=SimpleNamespace(store=SimpleNamespace(root=Path(tmp),project=lambda p:None))
            prepared=prepare(svc,'p',self.snapshot());folder=Path(tmp)/'report_packages/p'/prepared['package_id']
            html='<table><tr><th>行标识</th>'+('<th>来源说明</th>'*7)+'</tr>'
            html+=''.join('<tr><td>ROW-'+str(i).zfill(3)+'</td>'+('<td>'+'变动原因与适用范围。'*14+'</td>')*7+'</tr>' for i in range(70))
            html+='</table><p>END-OF-REPORT</p>'
            start=time.monotonic();finish(svc,'p',prepared['package_id'],html)
            self.assertLess(time.monotonic()-start,30)
            with fitz.open(folder/'report.pdf') as pdf:
                text='\n'.join(page.get_text() for page in pdf)
                self.assertIn('END-OF-REPORT',text)
                for i in range(70):self.assertEqual(text.count('ROW-'+str(i).zfill(3)),3)

    def test_worker_timeout_cleans_partial_pdf_and_releases_conversation(self):
        with tempfile.TemporaryDirectory() as tmp:
            svc=SimpleNamespace(store=SimpleNamespace(root=Path(tmp),project=lambda p:None))
            prepared=prepare(svc,'p',self.snapshot());folder=Path(tmp)/'report_packages/p'/prepared['package_id']
            def timeout(*args,**kwargs):
                (folder/'.report.pdf').write_bytes(b'');raise subprocess.TimeoutExpired(args[0],kwargs['timeout'])
            with patch('service.report_package.subprocess.run',side_effect=timeout):
                with self.assertRaisesRegex(RuntimeError,'已结束渲染进程'):finish(svc,'p',prepared['package_id'],'<p>test</p>')
            self.assertFalse((folder/'report.pdf').exists());self.assertFalse((folder/'.report.pdf').exists())
            self.assertTrue((folder/'workpapers.xlsx').is_file())
            finish(svc,'p',prepared['package_id'],'<p>Retry completes</p>')
            with fitz.open(folder/'report.pdf') as pdf:self.assertEqual(len(pdf),1)

    def test_exact_cells_unknowns_and_provenance_survive_workbook_and_zip(self):
        with tempfile.TemporaryDirectory() as tmp:
            report=compile_report({'tables':{'income':{'columns':{'key':{},'value':{'unit':'百万元','currency':'CNY'}},'rows':[
                {'key':'revenue','value':{'value':'12.1234567890123456789','unit':'百万元','currency':'CNY','source':{'document_id':'d','page':7}}},
                {'key':'=1+1','value':{'value':None,'unit':'百万元','currency':'CNY','source':{'missing_reason':'未披露'}}}]}}})
            user_source={'kind':'user_provided','messageId':'question','sourceId':'user:question','title':'用户提供','quote':'收入由用户提供。'}
            snapshot={'project_id':'p','conversation_id':'c','title':'测试报告','turns':[{'id':'m','created_at':'2026-10-08','body':'【观点】关注现金转换。','sources':[user_source],'tables':report['tables']}]}
            svc=SimpleNamespace(store=SimpleNamespace(root=Path(tmp),project=lambda p:None))
            prepared=prepare(svc,'p',snapshot);folder=Path(tmp)/'report_packages/p'/prepared['package_id']
            book=load_workbook(folder/'workpapers.xlsx');sheet=book['指标']
            self.assertEqual(json.loads(book['来源']['E4'].value),user_source)
            row=next(row for row in sheet if row[0].value=='revenue');self.assertEqual(row[1].value,'12.1234567890123456789')
            self.assertIn('"page": 7',row[1].comment.text)
            blank=next(row for row in sheet if row[0].value=='=1+1');self.assertEqual(blank[0].data_type,'s');self.assertIsNone(blank[1].value)
            self.assertIn('未披露',blank[1].comment.text)
            mixed=compile_report({'tables':{'mixed':{'columns':{'key':{},'value':{}},'rows':[
                {'key':'cash','value':{'value':'12','unit':'元','currency':'CNY'}},
                {'key':'shares','value':{'value':'100','unit':'股'}}]}}})
            from service.report_package import workbook
            workbook({**snapshot,'turns':[{**snapshot['turns'][0],'tables':mixed['tables']}]},folder/'mixed.xlsx')
            mixed_sheet=load_workbook(folder/'mixed.xlsx')['指标']
            self.assertEqual(mixed_sheet['C6'].value,'CNY 元');self.assertEqual(mixed_sheet['C7'].value,'股')
            self.assertEqual(json.loads((folder/'snapshot.json').read_text('utf8')),snapshot)
            exported=finish(svc,'p',prepared['package_id'],'<h1>财报分析报告</h1><p>收入 12.12 CNY 百万元</p>')
            self.assertTrue((Path(tmp)/exported['relative_path']).is_file())
            with fitz.open(folder/'report.pdf') as pdf:self.assertIn('12.12',pdf[0].get_text())
            self.assertEqual(section_for('new_kind',{'report_section':'anomaly_signals'}),4)
            with self.assertRaisesRegex(ValueError,'项目'):prepare(svc,'other',snapshot)
            with self.assertRaisesRegex(ValueError,'标识'):finish(svc,'p','../outside','')

if __name__=='__main__':unittest.main()

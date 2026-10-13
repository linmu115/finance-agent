"""Compare every collected attachment with its anonymous representation, offline."""
import argparse
import csv
import hashlib
import io
import json
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import fitz
import openpyxl
from blind import Redactor, TEXT, read, dump, tools_in, output


def verify(study, packet):
    study, packet = Path(study), output(packet)
    checks = []

    def check(raw, dest, ext, red, label, source_name):
        rec = {'label': label, 'original': source_name,
               'anonymous': str(dest.relative_to(packet)), 'format': ext,
               'original_sha256': hashlib.sha256(raw).hexdigest(),
               'original_rows': None, 'anonymous_rows': None,
               'original_columns': None, 'anonymous_columns': None,
               'headers_equal_except_identifiers': None,
               'numeric_cells': 0, 'identity_strings_changed': 0,
               'mismatches': [], 'status': 'passed'}

        def scalar(a, b, where):
            if type(a) in (int, float):
                rec['numeric_cells'] += 1
            elif isinstance(a, str):
                try:
                    float(a.strip().replace(',', '').removesuffix('%'))
                    rec['numeric_cells'] += 1
                except ValueError:
                    pass
            expected = red(a) if isinstance(a, str) else a
            if expected != a:
                rec['identity_strings_changed'] += 1
            if b != expected or (not isinstance(a, str) and type(b) != type(a)):
                rec['mismatches'].append({'at': where, 'reason': 'nonidentity_value_or_type_changed'})

        def tree(a, b, where='$'):
            if isinstance(a, dict):
                if not isinstance(b, dict) or set(b) != {red(k) for k in a}:
                    rec['mismatches'].append({'at': where, 'reason': 'object_keys_changed'}); return
                for k, v in a.items():
                    if k in {'profile', 'evaluation_profile', 'agent_profile'} and v in ('original_pi','prepared','financial'):
                        scalar('system_component', b[red(k)], where+'.'+k)
                        rec['identity_strings_changed'] += 1
                    else:
                        tree(v, b[red(k)], where+'.'+k)
            elif isinstance(a, list):
                if not isinstance(b, list) or len(a) != len(b):
                    rec['mismatches'].append({'at': where, 'reason': 'array_shape_changed'}); return
                for i, (x, y) in enumerate(zip(a, b)): tree(x, y, f'{where}[{i}]')
            else:
                scalar(a, b, where)

        if ext == '.csv':
            original = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'), newline='')))
            with dest.open(encoding='utf8', newline='') as h: anonymous = list(csv.reader(h))
            rec.update(original_rows=len(original), anonymous_rows=len(anonymous),
                       original_columns=dict(Counter(map(len, original))),
                       anonymous_columns=dict(Counter(map(len, anonymous))))
            rec['headers_equal_except_identifiers'] = bool(not original and not anonymous) or bool(original and anonymous and [red(c) for c in original[0]] == anonymous[0])
            if [len(r) for r in original] != [len(r) for r in anonymous]:
                rec['mismatches'].append({'reason': 'csv_row_or_column_counts_changed'})
            for i, (a, b) in enumerate(zip(original, anonymous)):
                for j, (x, y) in enumerate(zip(a, b)): scalar(x, y, f'row{i+1}/col{j+1}')
        elif ext == '.json':
            tree(json.loads(raw.decode('utf-8-sig')), read(dest))
        elif ext in TEXT:
            scalar(raw.decode('utf-8-sig').replace('\r\n', '\n'), dest.read_text('utf8'), 'full_text')
        elif ext in {'.png', '.jpg', '.jpeg'}:
            if raw != dest.read_bytes(): rec['mismatches'].append({'reason':'image_bytes_changed'})
        elif ext in {'.pdf', '.svg'}:
            doc = fitz.open(stream=raw, filetype=ext[1:])
            for i, page in enumerate(doc, 1):
                png = dest/(f'page-{i:03}.png' if ext == '.pdf' else 'rendered.png')
                expected = page.get_pixmap(matrix=fitz.Matrix(1.5,1.5), alpha=False)
                actual = fitz.Pixmap(str(png))
                if (actual.width,actual.height,actual.samples)!=(expected.width,expected.height,expected.samples):
                    rec['mismatches'].append({'at':f'page{i}','reason':'original_pixels_changed'})
                txt = dest/(f'page-{i:03}.txt' if ext == '.pdf' else 'visible-text.txt')
                scalar(page.get_text(), txt.read_text('utf8'), f'page{i}/text')
            rec['original_rows'] = rec['anonymous_rows'] = len(doc)
            if ext == '.svg': scalar(raw.decode('utf-8-sig'), (dest/'source.svg.txt').read_text('utf8'), 'svg_source')
            doc.close()
        elif ext == '.xlsx':
            formulas = openpyxl.load_workbook(io.BytesIO(raw),data_only=False)
            cached = openpyxl.load_workbook(io.BytesIO(raw),data_only=True)
            count = 0
            for i, sheet in enumerate(formulas,1):
                got = read(dest/f'sheet-{i:03}.json')
                scalar(sheet.title,got['sheet'],f'sheet{i}/name')
                expected=[]
                for row in sheet:
                    for cell in row:
                        if cell.value is None: continue
                        expected.append({'cell':cell.coordinate,'value_or_formula':cell.value,'cached_value':cached[sheet.title][cell.coordinate].value,'number_format':cell.number_format,'style_id':cell.style_id,'comment':cell.comment.text if cell.comment else None,'hidden_row':bool(sheet.row_dimensions[cell.row].hidden)})
                expected=json.loads(json.dumps(expected,ensure_ascii=False,default=str))
                tree(expected,got['cells'],f'sheet{i}/cells')
                tree([str(r) for r in sheet.merged_cells.ranges],got['merged_ranges'],f'sheet{i}/merged')
                scalar(sheet.sheet_state,got['state'],f'sheet{i}/state')
                count += len(expected)
            rec['original_rows']=rec['anonymous_rows']=count
            rec['original_columns']=rec['anonymous_columns']=len(formulas.sheetnames)
            visuals=read(dest/'document-review.json')['visual_assets']
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                assets=[n for n in z.namelist() if (n.startswith('xl/charts/') and n.endswith('.xml')) or (n.startswith('xl/media/') and n.lower().endswith(('.png','.jpg','.jpeg')))]
                if len(assets)!=len(visuals): rec['mismatches'].append({'reason':'xlsx_visual_asset_count_changed'})
                for original, anon in zip(assets,visuals):
                    data=z.read(original)
                    if original.endswith('.xml'):scalar(data.decode('utf8'),(dest/anon).read_text('utf8'),original)
                    elif data!=(dest/anon).read_bytes():rec['mismatches'].append({'at':original,'reason':'xlsx_image_changed'})
            formulas.close();cached.close()
        else:
            raise ValueError(ext)
        if rec['mismatches']: rec['status']='failed'
        checks.append(rec)

    def archive(raw, dest, replacements, label, source_name):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            members=[m for m in z.infolist() if not m.is_dir()]
            names={m.filename:f'member_{i:03}'+Path(m.filename).suffix for i,m in enumerate(members,1)}
            red=Redactor({**replacements,**names})
            for m in members:
                ext=Path(m.filename).suffix.lower();out=dest/names[m.filename]
                if ext=='.zip':archive(z.read(m),out.with_suffix('.zip.expanded'),{**replacements,**names},label,source_name+'!'+m.filename)
                else:check(z.read(m),out.with_suffix(ext+'.expanded') if ext in {'.pdf','.xlsx','.svg'} else out,ext,red,label,source_name+'!'+m.filename)

    primary=[]
    for case in read(packet/'private-map.json')['cases']:
        replacements=dict(tools_in(study));names={}
        for f in case['files']:names.setdefault(f['original']['name'],[]).append(f)
        for f in case['files']:
            a=f['original'];replacements[a['url']]=f['anonymous']
            for key in ['localFile','evaluatorLocalFile']:
                if a.get(key):replacements[a[key]]=f['anonymous']
        for name, fs in names.items():
            replacements[name]=fs[0]['anonymous'] if len(fs)==1 else fs[0]['anonymous'].split('_v')[0]+'_[ambiguous_version]'+Path(name).suffix
        red=Redactor(replacements);red.ids.update(case['replacements'])
        for f in case['files']:
            a=f['original'];local=a.get('localFile') or a.get('evaluatorLocalFile')
            if not local: raise ValueError('Unavailable collected attachment: '+case['label'])
            src=study/(case['label']+'-artifacts')/local
            ext=Path(f['anonymous']).suffix;dest=packet/'packet'/case['case_id']/f['anonymous']
            before=len(checks)
            if ext=='.zip':archive(src.read_bytes(),dest.with_suffix('.zip.expanded'),replacements,case['label'],a['name'])
            else:check(src.read_bytes(),dest.with_suffix(ext+'.expanded') if ext in {'.pdf','.xlsx','.svg'} else dest,ext,red,case['label'],a['name'])
            primary.append({'label':case['label'],'original':a['name'],'anonymous':f['anonymous'],'format':ext,'constituent_files':len(checks)-before,'numeric_cells':sum(x['numeric_cells'] for x in checks[before:]),'status':'passed' if all(x['status']=='passed' for x in checks[before:]) else 'failed'})
        delivery=read(study/(case['label']+'-deliverables.json'))
        for i, script in enumerate(delivery.get('evaluatorScripts',[]),1):
            src=study/(case['label']+'-artifacts')/script['file']
            dest=packet/'packet'/case['case_id']/f'code_{i:03}.py'
            check(src.read_bytes(),dest,'.py',red,case['label'],'evaluator_code:'+script['file'])
    result={'at':datetime.now(timezone.utc).isoformat(),'primary_attachments':len(primary),'expanded_file_checks':len(checks),'all_passed':all(x['status']=='passed' for x in checks),'method':'CSV row/column/header and every cell; JSON every typed leaf/key; XLSX every populated cell/formula/cache/style/comment and embedded image/chart; original PNG bytes and PDF/SVG page pixels; full text after source/identity-only substitutions. No candidate code executed.','checks':checks,'primary':primary}
    dump(packet/'attachment-verification.json',result)
    with (packet/'attachment-verification.csv').open('w',encoding='utf-8-sig',newline='') as h:
        writer=csv.DictWriter(h,fieldnames=['label','original','anonymous','format','original_rows','anonymous_rows','original_columns','anonymous_columns','headers_equal_except_identifiers','numeric_cells','identity_strings_changed','status','mismatch_count'])
        writer.writeheader()
        for c in checks:writer.writerow({k:len(c['mismatches']) if k=='mismatch_count' else c[k] for k in writer.fieldnames})
    print(json.dumps({k:result[k] for k in ['primary_attachments','expanded_file_checks','all_passed']}))
    if not result['all_passed']:raise SystemExit(1)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('study');parser.add_argument('packet');a=parser.parse_args();verify(a.study,a.packet)
